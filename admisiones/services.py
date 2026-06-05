import os
import re
import pythoncom

from django.conf import settings
from django.contrib.staticfiles import finders
from django.core.mail import EmailMessage
from django.template.loader import render_to_string
from django.utils import timezone
from django.utils.html import escape

from docx import Document
from docx2pdf import convert
from xhtml2pdf import pisa


# =========================================================
# UTILIDADES GENERALES
# =========================================================

def link_callback(uri, rel):
    """
    Convierte rutas de static/media a rutas absolutas para xhtml2pdf.
    """
    if uri.startswith(settings.STATIC_URL):
        path = finders.find(uri.replace(settings.STATIC_URL, ""))
        if path:
            return path

    if uri.startswith(settings.MEDIA_URL):
        return os.path.join(settings.MEDIA_ROOT, uri.replace(settings.MEDIA_URL, ""))

    posible_static = finders.find(uri)
    if posible_static:
        return posible_static

    return os.path.join(settings.BASE_DIR, uri)


def _replace(doc: Document, mapping: dict):
    """
    Busca y reemplaza texto en párrafos y tablas de un documento Word.
    """
    for p in doc.paragraphs:
        for k, v in mapping.items():
            if k in p.text:
                for run in p.runs:
                    if k in run.text:
                        run.text = run.text.replace(k, str(v))

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for k, v in mapping.items():
                        if k in p.text:
                            for run in p.runs:
                                if k in run.text:
                                    run.text = run.text.replace(k, str(v))


def _limpiar_nombre_archivo(valor):
    valor = str(valor or "").strip()
    valor = re.sub(r"[^A-Za-z0-9ÁÉÍÓÚáéíóúÑñ_-]+", "_", valor)
    return valor.strip("_") or "SIN_NOMBRE"


def _fecha_corta(valor):
    if not valor:
        return "No registrada"

    try:
        return valor.strftime("%d/%m/%Y")
    except Exception:
        return str(valor)


def _fecha_larga(valor=None):
    valor = valor or timezone.now()
    meses = {
        1: "enero",
        2: "febrero",
        3: "marzo",
        4: "abril",
        5: "mayo",
        6: "junio",
        7: "julio",
        8: "agosto",
        9: "septiembre",
        10: "octubre",
        11: "noviembre",
        12: "diciembre",
    }

    return f"{valor.day:02d} de {meses.get(valor.month, '')} de {valor.year}"


def _valor(valor, default="No registrado"):
    if valor is None:
        return default

    valor = str(valor).strip()
    return valor if valor else default


def _nombre_tutor(solicitud):
    partes = [
        getattr(solicitud, "nombres_tutor", ""),
        getattr(solicitud, "apellido_paterno_tutor", ""),
        getattr(solicitud, "apellido_materno_tutor", ""),
    ]
    return " ".join([p for p in partes if p]).strip() or "No registrado"


def _nombre_menor(solicitud):
    partes = [
        getattr(solicitud, "nombres_menor", ""),
        getattr(solicitud, "apellido_paterno_menor", ""),
        getattr(solicitud, "apellido_materno_menor", ""),
    ]
    return " ".join([p for p in partes if p]).strip() or "No registrado"


def _domicilio_tutor(solicitud):
    partes = [
        getattr(solicitud, "calle", ""),
        getattr(solicitud, "numero_exterior", ""),
        getattr(solicitud, "numero_interior", ""),
        getattr(solicitud, "codigo_postal", ""),
        getattr(solicitud, "ciudad", ""),
        getattr(solicitud, "estado_domicilio", ""),
    ]
    partes_limpias = [str(p).strip() for p in partes if p]
    return ", ".join(partes_limpias) if partes_limpias else "No registrado"


def _ruta_logo(nombre_archivo):
    """
    Busca logo_unach.png o logo_eiu.png en static/img/.
    También intenta en static/ directo por si estuvieran ahí.
    """
    posibles = [
        f"img/{nombre_archivo}",
        nombre_archivo,
    ]

    for ruta in posibles:
        encontrada = finders.find(ruta)
        if encontrada and os.path.exists(encontrada):
            return encontrada.replace("\\", "/")

    rutas_fallback = [
        os.path.join(settings.BASE_DIR, "static", "img", nombre_archivo),
        os.path.join(settings.BASE_DIR, "static", nombre_archivo),
    ]

    for ruta in rutas_fallback:
        if os.path.exists(ruta):
            return ruta.replace("\\", "/")

    return ""


def _img_logo_html(ruta, texto_alt):
    if ruta:
        return f'<img src="{ruta}" class="logo" alt="{texto_alt}">'
    return f'<div class="logo-text">{texto_alt}</div>'




# =========================================================
# REGLAS GENERALES DE SEMESTRE / OPERACIÓN DEL SISTEMA
# =========================================================

def _estado_semestre_texto(semestre):
    """
    Devuelve el estado del semestre como texto normalizado.
    Soporta choices tipo TextChoices y valores string normales.
    """
    if not semestre:
        return ""

    return str(getattr(semestre, "estado", "") or "").strip().upper()


def _valor_estado_semestre(nombre_estado, default):
    """
    Obtiene valores desde Semestre.EstadoSemestre si existe.
    Así funciona aunque el modelo use TextChoices o strings simples.
    """
    try:
        from admisiones.models import Semestre

        estado_cls = getattr(Semestre, "EstadoSemestre", None)
        if estado_cls and hasattr(estado_cls, nombre_estado):
            return getattr(estado_cls, nombre_estado)
    except Exception:
        pass

    return default


def _query_semestres_base(sede=None):
    """
    Query base para semestres.
    Si en el futuro se filtra por sede, ya queda preparado.
    """
    from admisiones.models import Semestre

    qs = Semestre.objects.all()

    if sede is not None:
        # Soporta enviar objeto sede o id.
        if hasattr(sede, "id"):
            qs = qs.filter(sede=sede)
        else:
            qs = qs.filter(sede_id=sede)

    return qs


def obtener_semestre_activo_operativo(sede=None):
    """
    Semestre que permite operación normal del sistema.

    Regla:
    - ACTIVO / INICIADO / ABIERTO = permite operación.
    - CONFIGURACIÓN / CERRADO / sin semestre = solo consulta para roles operativos.
    """
    activo = _valor_estado_semestre("ACTIVO", "ACTIVO")

    estados_activos = {
        str(activo),
        "ACTIVO",
        "INICIADO",
        "ABIERTO",
    }

    return (
        _query_semestres_base(sede=sede)
        .filter(estado__in=list(estados_activos))
        .order_by("-fecha_inicio", "-id")
        .first()
    )


def obtener_ultimo_semestre_registrado(sede=None):
    """
    Último semestre registrado, aunque esté cerrado o en configuración.
    Sirve para mostrar historial/mensajes de modo solo consulta.
    """
    return (
        _query_semestres_base(sede=sede)
        .order_by("-fecha_inicio", "-id")
        .first()
    )


def hay_semestre_activo_operativo(sede=None):
    """
    True si existe un semestre ACTIVO/INICIADO para operar.
    """
    return obtener_semestre_activo_operativo(sede=sede) is not None


def semestre_esta_activo(semestre):
    """
    True si un semestre está en estado operativo.
    """
    if not semestre:
        return False

    estado = _estado_semestre_texto(semestre)
    return estado in {"ACTIVO", "INICIADO", "ABIERTO"}


def semestre_esta_cerrado(semestre):
    """
    True si el semestre ya fue cerrado.
    """
    if not semestre:
        return False

    cerrado = str(_valor_estado_semestre("CERRADO", "CERRADO")).upper()
    estado = _estado_semestre_texto(semestre)

    return estado in {cerrado, "CERRADO", "FINALIZADO"}


def semestre_en_configuracion(semestre):
    """
    True si el semestre todavía se está preparando.
    """
    if not semestre:
        return False

    configuracion = str(_valor_estado_semestre("CONFIGURACION", "CONFIGURACION")).upper()
    estado = _estado_semestre_texto(semestre)

    return estado in {
        configuracion,
        "CONFIGURACION",
        "CONFIGURACIÓN",
        "BORRADOR",
        "PENDIENTE",
    }


def semestre_permite_operacion(semestre):
    """
    Regla para roles operativos:
    solo pueden operar si el semestre está activo/iniciado.
    """
    return semestre_esta_activo(semestre)


def semestre_permite_edicion_admin(semestre):
    """
    Regla para Administrador del Sistema:
    puede configurar o ajustar semestres mientras NO estén cerrados.
    Un semestre cerrado queda solo lectura.
    """
    if not semestre:
        return True

    return not semestre_esta_cerrado(semestre)


def usuario_es_admin_sistema(user):
    """
    Valida Administrador del Sistema de forma flexible.
    Se usa para permitir que el administrador configure semestres aunque no exista
    semestre activo.
    """
    if not user or not getattr(user, "is_authenticated", False):
        return False

    if getattr(user, "is_superuser", False):
        return True

    if getattr(user, "tipo_persona", "") == "ADMIN":
        return True

    try:
        return user.groups.filter(name__iexact="Administrador del Sistema").exists()
    except Exception:
        return False


def contexto_semestre_operativo(sede=None):
    """
    Contexto estándar para templates y paneles.

    Variables principales:
    - semestre_activo_operativo
    - hay_semestre_activo_operativo
    - modo_solo_consulta
    - aviso_semestre_operativo
    - ultimo_semestre_registrado

    Regla:
    - Si hay semestre ACTIVO: operación normal.
    - Si no hay semestre ACTIVO: roles operativos solo consultan.
    - Si el último semestre está CERRADO: se informa que está cerrado.
    - Si el último semestre está en CONFIGURACIÓN: se informa que aún no se inicia.
    """
    semestre_activo = obtener_semestre_activo_operativo(sede=sede)
    ultimo_semestre = obtener_ultimo_semestre_registrado(sede=sede)

    hay_activo = bool(semestre_activo)
    modo_solo_consulta = not hay_activo

    if hay_activo:
        aviso = ""
        estado_operativo = "ACTIVO"
    elif ultimo_semestre and semestre_esta_cerrado(ultimo_semestre):
        aviso = (
            "El último semestre registrado está cerrado. Puedes consultar información histórica, "
            "pero no puedes operar, agregar, modificar ni eliminar registros hasta que el Administrador "
            "inicie un nuevo semestre."
        )
        estado_operativo = "CERRADO_SIN_ACTIVO"
    elif ultimo_semestre and semestre_en_configuracion(ultimo_semestre):
        aviso = (
            "Existe un semestre en configuración, pero todavía no ha sido iniciado. "
            "Puedes consultar información existente, pero la operación se habilitará hasta que "
            "el Administrador inicie el semestre."
        )
        estado_operativo = "CONFIGURACION_SIN_ACTIVO"
    else:
        aviso = (
            "No hay semestre activo/iniciado. Puedes iniciar sesión y consultar información existente, "
            "pero no puedes operar, agregar, modificar ni eliminar registros hasta que el Administrador "
            "cree e inicie un nuevo semestre."
        )
        estado_operativo = "SIN_SEMESTRE"

    return {
        "semestre_activo_operativo": semestre_activo,
        "hay_semestre_activo_operativo": hay_activo,
        "modo_solo_consulta": modo_solo_consulta,
        "aviso_semestre_operativo": aviso,
        "ultimo_semestre_registrado": ultimo_semestre,
        "ultimo_semestre_cerrado": bool(ultimo_semestre and semestre_esta_cerrado(ultimo_semestre)),
        "ultimo_semestre_en_configuracion": bool(ultimo_semestre and semestre_en_configuracion(ultimo_semestre)),
        "estado_operativo_sistema": estado_operativo,
    }


# =========================================================
# PDF FO-301-02-01 DESDE WORD
# =========================================================

def generar_pdf_fo301(solicitud):
    template_path = os.path.join(settings.BASE_DIR, "plantillas", "FO-301-02-01_template.docx")
    doc = Document(template_path)

    mapping = {
        "{{identificador}}": f"EI-{solicitud.created_at.year}-{str(solicitud.folio).zfill(4)}",
        "{{fecha}}": solicitud.created_at.strftime("%d/%m/%Y"),

        "{{nombre_beneficiario}}": solicitud.nombre_tutor_completo(),
        "{{matricula}}": getattr(solicitud, "matricula", ""),
        "{{programa_educativo}}": getattr(solicitud, "programa_educativo", ""),
        "{{unidad_academica}}": getattr(solicitud, "unidad_academica", ""),
        "{{correo_beneficiario}}": getattr(solicitud, "correo", ""),
        "{{domicilio_beneficiario}}": (
            f"{getattr(solicitud, 'calle', '')} "
            f"{getattr(solicitud, 'numero_exterior', '')}, "
            f"{getattr(solicitud, 'ciudad', '')}, "
            f"{getattr(solicitud, 'estado_domicilio', '')}"
        ),
        "{{celular_beneficiario}}": getattr(solicitud, "celular", ""),
        "{{parentesco_con_menor}}": getattr(solicitud, "parentesco", ""),

        "{{nombre_menor}}": getattr(solicitud, "nombres_menor", ""),
        "{{fecha_nacimiento}}": (
            solicitud.fecha_nacimiento_menor.strftime("%d/%m/%Y")
            if getattr(solicitud, "fecha_nacimiento_menor", None)
            else ""
        ),
        "{{curp_menor}}": getattr(solicitud, "curp_menor", ""),
    }

    _replace(doc, mapping)

    out_docx_dir = os.path.join(settings.MEDIA_ROOT, "solicitudes", "docx")
    out_pdf_dir = os.path.join(settings.MEDIA_ROOT, "solicitudes", "pdf")
    os.makedirs(out_docx_dir, exist_ok=True)
    os.makedirs(out_pdf_dir, exist_ok=True)

    out_docx_path = os.path.join(out_docx_dir, f"Solicitud_{solicitud.folio}.docx")
    out_pdf_path = os.path.join(out_pdf_dir, f"Solicitud_{solicitud.folio}.pdf")

    doc.save(out_docx_path)

    try:
        pythoncom.CoInitialize()
        convert(out_docx_path, out_pdf_path)

        if os.path.exists(out_docx_path):
            os.remove(out_docx_path)

    except Exception as e:
        print(f"Error en la conversión de DOCX a PDF: {e}")

    return out_pdf_path


# =========================================================
# DICTAMEN DE INGRESO — ACEPTACIÓN / RECHAZO
# SIN TEMPLATE HTML EXTERNO
# =========================================================

def generar_pdf_dictamen(solicitud):
    """
    Genera PDF detallado de aceptación o rechazo de solicitud de ingreso.
    No usa template HTML externo.
    Guarda en:
    media/admisiones/dictamenes_ingreso/
    """
    estado_actual = str(getattr(solicitud, "estado", "")).upper()

    if estado_actual == "APROBADA":
        dictamen = "ACEPTADA"
        titulo = "Carta de aceptación de solicitud de ingreso"
        color_principal = "#0f766e"
        color_fondo = "#ecfdf5"
        color_borde = "#14b8a6"
        texto_resolucion = (
            "Después de revisar la información proporcionada en la solicitud de ingreso, "
            "se informa que la solicitud ha sido aceptada para continuar con el proceso de inscripción "
            "a la Estancia Infantil Universitaria."
        )
        texto_siguiente_paso = (
            "El beneficiario deberá ingresar al sistema con su cuenta institucional para completar "
            "la solicitud de inscripción, cargar la documentación requerida, aceptar reglamentos "
            "y enviar el expediente para revisión."
        )
        resultado_visible = "SOLICITUD ACEPTADA"

    else:
        dictamen = "RECHAZADA"
        titulo = "Carta de rechazo de solicitud de ingreso"
        color_principal = "#991b1b"
        color_fondo = "#fef2f2"
        color_borde = "#ef4444"
        texto_resolucion = (
            "Después de revisar la información proporcionada en la solicitud de ingreso, "
            "se informa que la solicitud no fue aceptada para continuar con el proceso de inscripción "
            "en este momento."
        )
        texto_siguiente_paso = (
            "El beneficiario podrá comunicarse con el área correspondiente para recibir orientación "
            "sobre los motivos del dictamen o sobre los pasos a seguir, en caso de que aplique."
        )
        resultado_visible = "SOLICITUD RECHAZADA"

    comentarios = (
        getattr(solicitud, "comentario_respuesta", None)
        or "Cumple con los requisitos establecidos por la Estancia Infantil Universitaria."
        if dictamen == "ACEPTADA"
        else getattr(solicitud, "comentario_respuesta", None)
        or "No se especificó un motivo adicional."
    )

    folio = _valor(getattr(solicitud, "folio", "SIN_FOLIO"))
    identificador = f"EI-{timezone.now().year}-{str(folio).zfill(4)}" if str(folio).isdigit() else f"EI-{folio}"

    ruta_logo_unach = _ruta_logo("logo_unach.png")
    ruta_logo_eiu = _ruta_logo("logo_eiu.png")

    logo_unach_html = _img_logo_html(ruta_logo_unach, "UNACH")
    logo_eiu_html = _img_logo_html(ruta_logo_eiu, "EIU")

    nombre_tutor = _nombre_tutor(solicitud)
    nombre_menor = _nombre_menor(solicitud)

    fecha_solicitud = _fecha_corta(getattr(solicitud, "created_at", None))
    fecha_dictamen = _fecha_larga(timezone.now())

    matricula = _valor(getattr(solicitud, "matricula", ""))
    correo = _valor(getattr(solicitud, "correo", ""))
    celular = _valor(getattr(solicitud, "celular", ""))
    unidad = _valor(getattr(solicitud, "unidad_academica", ""))
    programa = _valor(getattr(solicitud, "programa_educativo", ""))
    parentesco = _valor(getattr(solicitud, "parentesco", ""))
    domicilio = _domicilio_tutor(solicitud)
    curp_menor = _valor(getattr(solicitud, "curp_menor", ""))
    fecha_nacimiento = _fecha_corta(getattr(solicitud, "fecha_nacimiento_menor", None))

    dictaminado_por = getattr(solicitud, "dictaminado_por", None)
    if dictaminado_por:
        nombre_dictaminador = (
            dictaminado_por.get_full_name()
            if hasattr(dictaminado_por, "get_full_name")
            else str(dictaminado_por)
        )
        if not nombre_dictaminador:
            nombre_dictaminador = getattr(dictaminado_por, "email", "Personal autorizado")
    else:
        nombre_dictaminador = "Personal autorizado"

    html = f"""
    <!DOCTYPE html>
    <html lang="es">
    <head>
        <meta charset="UTF-8">
        <style>
            @page {{
                size: letter;
                margin: 1.4cm 1.4cm 1.6cm 1.4cm;
            }}

            body {{
                font-family: Arial, Helvetica, sans-serif;
                font-size: 10.5px;
                color: #1f2937;
                line-height: 1.35;
            }}

            .header-table {{
                width: 100%;
                border-bottom: 3px solid {color_principal};
                padding-bottom: 8px;
                margin-bottom: 14px;
            }}

            .header-table td {{
                vertical-align: middle;
            }}

            .logo-box {{
                width: 22%;
                text-align: center;
            }}

            .logo {{
                max-height: 58px;
                max-width: 120px;
            }}

            .logo-text {{
                border: 1px solid #d1d5db;
                padding: 10px;
                font-size: 11px;
                font-weight: bold;
                color: #374151;
            }}

            .title-box {{
                width: 56%;
                text-align: center;
            }}

            .institution {{
                font-size: 14px;
                font-weight: bold;
                text-transform: uppercase;
                color: #111827;
                margin-bottom: 2px;
            }}

            .subtitle {{
                font-size: 11px;
                color: #374151;
                margin-bottom: 3px;
            }}

            .document-title {{
                font-size: 13px;
                font-weight: bold;
                color: {color_principal};
                text-transform: uppercase;
            }}

            .meta-table {{
                width: 100%;
                border-collapse: collapse;
                margin-bottom: 12px;
            }}

            .meta-table td {{
                border: 1px solid #d1d5db;
                padding: 5px;
            }}

            .label {{
                font-weight: bold;
                background: #f3f4f6;
                width: 24%;
            }}

            .status-box {{
                border: 2px solid {color_borde};
                background: {color_fondo};
                padding: 12px;
                text-align: center;
                margin: 12px 0 14px 0;
            }}

            .status-title {{
                font-size: 18px;
                font-weight: bold;
                color: {color_principal};
                letter-spacing: 0.5px;
                margin-bottom: 4px;
            }}

            .section-title {{
                background: {color_principal};
                color: #ffffff;
                font-weight: bold;
                padding: 6px;
                margin-top: 10px;
                margin-bottom: 0;
                font-size: 11px;
                text-transform: uppercase;
            }}

            .data-table {{
                width: 100%;
                border-collapse: collapse;
                margin-bottom: 8px;
            }}

            .data-table td {{
                border: 1px solid #d1d5db;
                padding: 5px;
                vertical-align: top;
            }}

            .paragraph-box {{
                border: 1px solid #d1d5db;
                padding: 8px;
                text-align: justify;
                margin-bottom: 8px;
            }}

            .comment-box {{
                border: 1px solid {color_borde};
                padding: 8px;
                background: #fafafa;
                min-height: 45px;
                text-align: justify;
            }}

            .note {{
                font-size: 9.5px;
                color: #4b5563;
                text-align: justify;
                border-top: 1px solid #d1d5db;
                padding-top: 7px;
                margin-top: 10px;
            }}

            .signature-table {{
                width: 100%;
                margin-top: 34px;
                text-align: center;
            }}

            .signature-line {{
                border-top: 1px solid #111827;
                padding-top: 5px;
                width: 75%;
                margin: 0 auto;
            }}

            .footer {{
                position: fixed;
                bottom: 0.35cm;
                left: 0;
                right: 0;
                text-align: center;
                font-size: 8.5px;
                color: #6b7280;
                border-top: 1px solid #d1d5db;
                padding-top: 4px;
            }}
        </style>
    </head>
    <body>
        <table class="header-table">
            <tr>
                <td class="logo-box">{logo_unach_html}</td>
                <td class="title-box">
                    <div class="institution">Universidad Autónoma de Chiapas</div>
                    <div class="subtitle">Estancia Infantil Universitaria</div>
                    <div class="document-title">{escape(titulo)}</div>
                </td>
                <td class="logo-box">{logo_eiu_html}</td>
            </tr>
        </table>

        <table class="meta-table">
            <tr>
                <td class="label">Folio</td>
                <td>{escape(str(folio))}</td>
                <td class="label">Identificador</td>
                <td>{escape(identificador)}</td>
            </tr>
            <tr>
                <td class="label">Fecha de solicitud</td>
                <td>{escape(fecha_solicitud)}</td>
                <td class="label">Fecha de dictamen</td>
                <td>{escape(fecha_dictamen)}</td>
            </tr>
        </table>

        <div class="status-box">
            <div class="status-title">{escape(resultado_visible)}</div>
            <div>Dictamen emitido por la Estancia Infantil Universitaria</div>
        </div>

        <div class="section-title">Datos del beneficiario</div>
        <table class="data-table">
            <tr>
                <td class="label">Nombre</td>
                <td>{escape(nombre_tutor)}</td>
                <td class="label">Matrícula</td>
                <td>{escape(matricula)}</td>
            </tr>
            <tr>
                <td class="label">Correo</td>
                <td>{escape(correo)}</td>
                <td class="label">Celular</td>
                <td>{escape(celular)}</td>
            </tr>
            <tr>
                <td class="label">Unidad académica</td>
                <td>{escape(unidad)}</td>
                <td class="label">Programa educativo</td>
                <td>{escape(programa)}</td>
            </tr>
            <tr>
                <td class="label">Parentesco</td>
                <td>{escape(parentesco)}</td>
                <td class="label">Domicilio</td>
                <td>{escape(domicilio)}</td>
            </tr>
        </table>

        <div class="section-title">Datos del menor</div>
        <table class="data-table">
            <tr>
                <td class="label">Nombre del menor</td>
                <td>{escape(nombre_menor)}</td>
            </tr>
            <tr>
                <td class="label">CURP</td>
                <td>{escape(curp_menor)}</td>
            </tr>
            <tr>
                <td class="label">Fecha de nacimiento</td>
                <td>{escape(fecha_nacimiento)}</td>
            </tr>
        </table>

        <div class="section-title">Resolución institucional</div>
        <div class="paragraph-box">
            {escape(texto_resolucion)}
        </div>

        <div class="section-title">Observaciones del dictamen</div>
        <div class="comment-box">
            {escape(str(comentarios))}
        </div>

        <div class="section-title">Indicaciones</div>
        <div class="paragraph-box">
            {escape(texto_siguiente_paso)}
        </div>

        <div class="note">
            Este documento se genera de manera electrónica como constancia del dictamen de la solicitud de ingreso.
            La aceptación de la solicitud de ingreso no sustituye la validación posterior del expediente de inscripción,
            documentación requerida, filtros institucionales o cualquier otro requisito establecido por la Estancia
            Infantil Universitaria.
        </div>

        <table class="signature-table">
            <tr>
                <td>
                    <div class="signature-line">
                        {escape(nombre_dictaminador)}<br>
                        Área responsable de admisión<br>
                        Estancia Infantil Universitaria
                    </div>
                </td>
            </tr>
        </table>

        <div class="footer">
            Universidad Autónoma de Chiapas · Estancia Infantil Universitaria · Documento generado automáticamente
        </div>
    </body>
    </html>
    """

    out_pdf_dir = os.path.join(settings.MEDIA_ROOT, "admisiones", "dictamenes_ingreso")
    os.makedirs(out_pdf_dir, exist_ok=True)

    nombre_archivo = f"Dictamen_Ingreso_{folio}_{dictamen}.pdf"
    nombre_archivo = _limpiar_nombre_archivo(nombre_archivo.replace(".pdf", "")) + ".pdf"
    out_pdf_path = os.path.join(out_pdf_dir, nombre_archivo)

    with open(out_pdf_path, "w+b") as result_file:
        pisa_status = pisa.CreatePDF(html, dest=result_file, link_callback=link_callback)

    if pisa_status.err:
        print(f"Error al generar el PDF del dictamen de ingreso {folio}")
        return None

    return out_pdf_path


def enviar_correo_dictamen_ingreso(solicitud, pdf_path=None):
    """
    Envía correo de aceptación o rechazo con el PDF adjunto.
    En desarrollo se imprime en consola si EMAIL_BACKEND está en console.
    """
    correo_destino = getattr(solicitud, "correo", "")

    if not correo_destino:
        return False, "La solicitud no tiene correo registrado."

    estado_actual = str(getattr(solicitud, "estado", "")).upper()
    nombre_tutor = _nombre_tutor(solicitud)
    folio = _valor(getattr(solicitud, "folio", "SIN_FOLIO"))

    if estado_actual == "APROBADA":
        asunto = f"Solicitud de ingreso aceptada - Folio {folio}"
        cuerpo = f"""
Estimado(a) {nombre_tutor}:

Por este medio se le informa que su solicitud de ingreso con folio {folio} ha sido ACEPTADA.

A partir de este momento podrá continuar con el proceso de inscripción en el sistema.
Para ingresar al portal utilice su correo institucional registrado. En caso de que se haya generado una contraseña temporal, esta será indicada por el área correspondiente.

Se adjunta el dictamen de aceptación en formato PDF.

Atentamente,
Estancia Infantil Universitaria
Universidad Autónoma de Chiapas
"""
    else:
        asunto = f"Solicitud de ingreso rechazada - Folio {folio}"
        motivo = (
            getattr(solicitud, "comentario_respuesta", None)
            or "No se especificó un motivo adicional."
        )

        cuerpo = f"""
Estimado(a) {nombre_tutor}:

Por este medio se le informa que su solicitud de ingreso con folio {folio} ha sido RECHAZADA.

Motivo u observaciones:
{motivo}

Se adjunta el dictamen correspondiente en formato PDF.

Atentamente,
Estancia Infantil Universitaria
Universidad Autónoma de Chiapas
"""

    email = EmailMessage(
        subject=asunto,
        body=cuerpo,
        from_email=settings.DEFAULT_FROM_EMAIL,
        to=[correo_destino],
    )

    if pdf_path and os.path.exists(pdf_path):
        email.attach_file(pdf_path)

    try:
        email.send(fail_silently=False)
        return True, "Correo enviado correctamente."
    except Exception as e:
        return False, f"No se pudo enviar el correo: {e}"



# =========================================================
# REPORTE FINAL DE INGRESO FO-301-02-18
# =========================================================

def _nombre_usuario(user, default="Personal autorizado"):
    """
    Obtiene el nombre visible de un usuario del sistema.
    """
    if not user:
        return default

    try:
        nombre = user.get_full_name()
    except Exception:
        nombre = ""

    if not nombre:
        nombre = getattr(user, "email", "") or getattr(user, "username", "")

    return nombre or default


def _check_pdf(condicion):
    """
    Devuelve una X para marcar visualmente una opción en el PDF.
    """
    return "X" if condicion else ""


def _motivos_reporte_final(inscripcion):
    """
    Obtiene los motivos/observaciones del Reporte Final.
    """
    motivos = getattr(inscripcion, "reporte_final_motivos", None)
    if motivos:
        return str(motivos).strip()

    observaciones = getattr(inscripcion, "observaciones", None)
    if observaciones:
        return str(observaciones).strip()

    return ""


def generar_pdf_reporte_final_ingreso(inscripcion):
    """
    Genera el PDF FO-301-02-18 Reporte Final de Ingreso.

    Guarda el archivo en:
    media/inscripciones/<matricula>/reportes_finales/

    Regresa la ruta absoluta del PDF generado.
    """
    solicitud = inscripcion.solicitud_ingreso

    folio = _valor(getattr(solicitud, "folio", "SIN_FOLIO"))
    matricula = _limpiar_nombre_archivo(
        getattr(solicitud, "matricula", "") or "SIN_MATRICULA"
    )
    nombre_menor = _nombre_menor(solicitud)
    nombre_tutor = _nombre_tutor(solicitud)

    fecha_base = (
        getattr(inscripcion, "directora_en", None)
        or getattr(inscripcion, "reporte_final_en", None)
        or timezone.now()
    )
    fecha_reporte = _fecha_larga(fecha_base)

    # La resolución final prioriza el dictamen de Dirección.
    # Si todavía no hay dictamen de Dirección, usa el dictamen de Trabajo Social.
    if getattr(inscripcion, "directora_aprobado", None) is not None:
        aprobado = bool(inscripcion.directora_aprobado)
    else:
        aprobado = bool(getattr(inscripcion, "reporte_final_aprobado", False))

    motivos = _motivos_reporte_final(inscripcion)
    if not motivos:
        motivos = (
            "El expediente cumple con los criterios requeridos para continuar el proceso."
            if aprobado
            else "No se especificaron motivos adicionales."
        )

    observaciones_directora = getattr(inscripcion, "directora_observaciones", "") or ""

    if observaciones_directora:
        observaciones_directora_html = f"""
        <div class="section-title">Observaciones de Dirección</div>
        <div class="motivos-box">{escape(observaciones_directora)}</div>
        """
    else:
        observaciones_directora_html = ""

    trabajadora_social = _nombre_usuario(
        getattr(inscripcion, "reporte_final_por", None),
        "Trabajadora Social",
    )
    directora = _nombre_usuario(
        getattr(inscripcion, "directora_por", None),
        "Directora",
    )

    ruta_logo_unach = _ruta_logo("logo_unach.png")
    ruta_logo_eiu = _ruta_logo("logo_eiu.png")

    logo_unach_html = _img_logo_html(ruta_logo_unach, "UNACH")
    logo_eiu_html = _img_logo_html(ruta_logo_eiu, "EIU")

    marca_aceptado = _check_pdf(aprobado)
    marca_rechazado = _check_pdf(not aprobado)

    try:
        estado_psicologia = solicitud.get_estado_psicologia_display()
    except Exception:
        estado_psicologia = "No registrado"

    try:
        estado_nutricion = solicitud.get_estado_nutricion_display()
    except Exception:
        estado_nutricion = "No registrado"

    try:
        estado_medico = solicitud.get_estado_medico_display()
    except Exception:
        estado_medico = "No registrado"

    resolucion_visible = "Aceptada" if aprobado else "Rechazada"

    html = f"""
    <!DOCTYPE html>
    <html lang="es">
    <head>
        <meta charset="UTF-8">
        <style>
            @page {{
                size: letter;
                margin: 1.3cm 1.5cm 1.4cm 1.5cm;
            }}

            body {{
                font-family: Arial, Helvetica, sans-serif;
                font-size: 11px;
                color: #111827;
                line-height: 1.45;
            }}

            .header-table {{
                width: 100%;
                border-collapse: collapse;
                margin-bottom: 10px;
            }}

            .header-table td {{
                vertical-align: middle;
            }}

            .logo-box {{
                width: 22%;
                text-align: center;
            }}

            .logo {{
                max-height: 62px;
                max-width: 126px;
            }}

            .logo-text {{
                border: 1px solid #d1d5db;
                padding: 10px;
                font-size: 11px;
                font-weight: bold;
                color: #374151;
            }}

            .title-box {{
                width: 56%;
                text-align: center;
            }}

            .institution {{
                font-size: 13px;
                font-weight: bold;
                text-transform: uppercase;
                color: #111827;
                margin-bottom: 2px;
            }}

            .subtitle {{
                font-size: 10.5px;
                color: #374151;
                margin-bottom: 2px;
                text-transform: uppercase;
            }}

            .document-title {{
                font-size: 15px;
                font-weight: bold;
                color: #111827;
                text-transform: uppercase;
                margin-top: 6px;
            }}

            .code {{
                font-size: 10px;
                font-weight: bold;
                color: #374151;
                margin-top: 2px;
            }}

            .top-line {{
                border-top: 2px solid #111827;
                margin: 4px 0 14px 0;
            }}

            .meta-table {{
                width: 100%;
                border-collapse: collapse;
                margin-bottom: 14px;
            }}

            .meta-table td {{
                border: 1px solid #d1d5db;
                padding: 5px;
            }}

            .label {{
                font-weight: bold;
                background: #f3f4f6;
                width: 20%;
            }}

            .date-box {{
                text-align: right;
                margin-bottom: 18px;
            }}

            .date-line {{
                display: inline-block;
                min-width: 230px;
                border-bottom: 1px solid #111827;
                text-align: center;
                padding: 0 4px 2px 4px;
            }}

            .paragraph {{
                text-align: justify;
                margin-bottom: 12px;
                font-size: 11.2px;
            }}

            .child-line {{
                display: inline-block;
                min-width: 330px;
                border-bottom: 1px solid #111827;
                text-align: center;
                font-weight: bold;
            }}

            .option-table {{
                width: 100%;
                border-collapse: collapse;
                margin: 10px 0 12px 0;
            }}

            .option-table td {{
                vertical-align: top;
                padding: 4px 3px;
            }}

            .checkbox {{
                display: inline-block;
                width: 14px;
                height: 14px;
                border: 1px solid #111827;
                text-align: center;
                line-height: 13px;
                font-weight: bold;
                margin-right: 4px;
            }}

            .motivos-box {{
                border: 1px solid #d1d5db;
                min-height: 84px;
                padding: 8px;
                text-align: justify;
                margin-top: 6px;
                margin-bottom: 12px;
                white-space: pre-line;
            }}

            .section-title {{
                background: #f3f4f6;
                border: 1px solid #d1d5db;
                font-weight: bold;
                padding: 5px;
                margin-top: 10px;
                margin-bottom: 0;
                text-transform: uppercase;
                font-size: 10.5px;
            }}

            .filters-table {{
                width: 100%;
                border-collapse: collapse;
                margin-bottom: 10px;
            }}

            .filters-table td {{
                border: 1px solid #d1d5db;
                padding: 5px;
            }}

            .signature-table {{
                width: 100%;
                margin-top: 45px;
                text-align: center;
            }}

            .signature-table td {{
                width: 50%;
                padding: 0 22px;
                vertical-align: bottom;
            }}

            .signature-line {{
                border-top: 1px solid #111827;
                padding-top: 6px;
                min-height: 38px;
            }}

            .footer {{
                position: fixed;
                bottom: 0.35cm;
                left: 0;
                right: 0;
                text-align: center;
                font-size: 8.5px;
                color: #6b7280;
                border-top: 1px solid #d1d5db;
                padding-top: 4px;
            }}
        </style>
    </head>
    <body>

        <table class="header-table">
            <tr>
                <td class="logo-box">{logo_unach_html}</td>
                <td class="title-box">
                    <div class="institution">Universidad Autónoma de Chiapas</div>
                    <div class="subtitle">Estancia Infantil Universitaria</div>
                    <div class="document-title">Reporte Final de Ingreso</div>
                </td>
                <td class="logo-box">
                    {logo_eiu_html}
                    <div class="code">FO-301-02-18</div>
                </td>
            </tr>
        </table>

        <div class="top-line"></div>

        <table class="meta-table">
            <tr>
                <td class="label">Folio</td>
                <td>{escape(str(folio))}</td>
                <td class="label">Matrícula</td>
                <td>{escape(str(matricula))}</td>
            </tr>
            <tr>
                <td class="label">Beneficiario</td>
                <td colspan="3">{escape(nombre_tutor)}</td>
            </tr>
        </table>

        <div class="date-box">
            Fecha:
            <span class="date-line">{escape(fecha_reporte)}</span>
        </div>

        <div class="paragraph">
            Derivado de la información proporcionada por las diversas áreas que participan
            en el proceso de inscripción al servicio que proporciona la Estancia Infantil UNACH,
            se determina que el menor
            <span class="child-line">{escape(nombre_menor)}</span>:
        </div>

        <table class="option-table">
            <tr>
                <td style="width: 18px;">
                    <span class="checkbox">{marca_aceptado}</span>
                </td>
                <td>
                    Se encuentra en condiciones para concluir el trámite y se autoriza la inscripción del mismo.
                </td>
            </tr>
            <tr>
                <td>
                    <span class="checkbox">{marca_rechazado}</span>
                </td>
                <td>
                    No se encuentra en condiciones para concluir el trámite, por lo que no se autoriza
                    la inscripción del mismo por los siguientes motivos:
                </td>
            </tr>
        </table>

        <div class="motivos-box">
            {escape(motivos)}
        </div>

        <div class="section-title">Resultado de filtros institucionales</div>
        <table class="filters-table">
            <tr>
                <td class="label">Psicología</td>
                <td>{escape(estado_psicologia)}</td>
                <td class="label">Nutrición</td>
                <td>{escape(estado_nutricion)}</td>
            </tr>
            <tr>
                <td class="label">Médico</td>
                <td>{escape(estado_medico)}</td>
                <td class="label">Resolución</td>
                <td>{escape(resolucion_visible)}</td>
            </tr>
        </table>

        {observaciones_directora_html}

        <table class="signature-table">
            <tr>
                <td>
                    <div class="signature-line">
                        {escape(trabajadora_social)}<br>
                        Nombre y Firma de la Trabajadora Social
                    </div>
                </td>
                <td>
                    <div class="signature-line">
                        {escape(directora)}<br>
                        Nombre y Firma de la Directora
                    </div>
                </td>
            </tr>
        </table>

        <div class="footer">
            Universidad Autónoma de Chiapas · Estancia Infantil Universitaria · FO-301-02-18
        </div>

    </body>
    </html>
    """

    out_pdf_dir = os.path.join(
        settings.MEDIA_ROOT,
        "inscripciones",
        matricula,
        "reportes_finales",
    )
    os.makedirs(out_pdf_dir, exist_ok=True)

    estado_archivo = "ACEPTADO" if aprobado else "RECHAZADO"
    nombre_archivo = f"Reporte_Final_Ingreso_{folio}_{estado_archivo}.pdf"
    nombre_archivo = _limpiar_nombre_archivo(nombre_archivo.replace(".pdf", "")) + ".pdf"
    out_pdf_path = os.path.join(out_pdf_dir, nombre_archivo)

    with open(out_pdf_path, "w+b") as result_file:
        pisa_status = pisa.CreatePDF(html, dest=result_file, link_callback=link_callback)

    if pisa_status.err:
        print(f"Error al generar Reporte Final de Ingreso FO-301-02-18. Folio: {folio}")
        return None

    return out_pdf_path


def enviar_correo_reporte_final_ingreso(inscripcion, pdf_path=None):
    """
    Envía al beneficiario el Reporte Final de Ingreso.
    Se usa cuando Dirección emite la resolución final.
    """
    solicitud = inscripcion.solicitud_ingreso
    correo_destino = getattr(solicitud, "correo", "")

    if not correo_destino:
        return False, "La solicitud no tiene correo registrado."

    folio = _valor(getattr(solicitud, "folio", "SIN_FOLIO"))
    nombre_tutor = _nombre_tutor(solicitud)
    nombre_menor = _nombre_menor(solicitud)

    if getattr(inscripcion, "directora_aprobado", None) is not None:
        aprobado = bool(inscripcion.directora_aprobado)
    else:
        aprobado = bool(getattr(inscripcion, "reporte_final_aprobado", False))

    if aprobado:
        asunto = f"Inscripción aprobada - Reporte Final de Ingreso - Folio {folio}"
        cuerpo = f"""
Estimado(a) {nombre_tutor}:

Por este medio se le informa que, después de la revisión institucional correspondiente, la inscripción del menor {nombre_menor} ha sido APROBADA.

Se adjunta el Reporte Final de Ingreso FO-301-02-18.

Atentamente,
Estancia Infantil Universitaria
Universidad Autónoma de Chiapas
"""
    else:
        motivos = _motivos_reporte_final(inscripcion) or "No se especificaron motivos adicionales."
        asunto = f"Inscripción no autorizada - Reporte Final de Ingreso - Folio {folio}"
        cuerpo = f"""
Estimado(a) {nombre_tutor}:

Por este medio se le informa que, después de la revisión institucional correspondiente, la inscripción del menor {nombre_menor} NO fue autorizada.

Motivos u observaciones:
{motivos}

Se adjunta el Reporte Final de Ingreso FO-301-02-18.

Atentamente,
Estancia Infantil Universitaria
Universidad Autónoma de Chiapas
"""

    email = EmailMessage(
        subject=asunto,
        body=cuerpo,
        from_email=settings.DEFAULT_FROM_EMAIL,
        to=[correo_destino],
    )

    if pdf_path and os.path.exists(pdf_path):
        email.attach_file(pdf_path)

    try:
        email.send(fail_silently=False)
        return True, "Correo enviado correctamente."
    except Exception as e:
        return False, f"No se pudo enviar el correo: {e}"


# =========================================================
# HOJA DE INSCRIPCIÓN — SE CONSERVA COMO LA TIENES
# =========================================================

def generar_pdf_inscripcion_lleno(inscripcion):
    """
    Toma los datos del expediente y genera un PDF con diseño de cuadrícula usando HTML.
    Este sí conserva el template admisiones/hoja_inscripcion_pdf.html.
    """
    s = inscripcion.solicitud_ingreso

    calle = getattr(s, "calle", "")
    num_ext = getattr(s, "numero_exterior", "")
    num_int = getattr(s, "numero_interior", "")
    cp = getattr(s, "codigo_postal", "")
    ciudad = getattr(s, "ciudad", "")
    estado_dom = getattr(s, "estado_domicilio", "")

    domicilio_tutor = (
        f"{calle} {num_ext} {num_int}, C.P. {cp}, {ciudad}, {estado_dom}"
        .replace("None", "")
        .replace("  ", " ")
        .strip(", ")
    )

    calle_p2 = getattr(inscripcion, "calle_segundo_padre", "")
    num_ext_p2 = getattr(inscripcion, "numero_exterior_segundo_padre", "")
    num_int_p2 = getattr(inscripcion, "numero_interior_segundo_padre", "")
    cp_p2 = getattr(inscripcion, "codigo_postal_segundo_padre", "")
    ciudad_p2 = getattr(inscripcion, "ciudad_segundo_padre", "")
    estado_p2 = getattr(inscripcion, "estado_domicilio_segundo_padre", "")

    domicilio_segundo = (
        f"{calle_p2} {num_ext_p2} {num_int_p2}, C.P. {cp_p2}, {ciudad_p2}, {estado_p2}"
        .replace("None", "")
        .replace("  ", " ")
        .strip(", ")
    )

    context = {
        "inscripcion": inscripcion,
        "s": s,
        "domicilio_tutor": domicilio_tutor,
        "domicilio_segundo": (
            domicilio_segundo
            if getattr(inscripcion, "nombres_segundo_padre", "")
            else "No aplica"
        ),
        "fecha_hoy": timezone.now(),
    }

    html = render_to_string("admisiones/hoja_inscripcion_pdf.html", context)

    out_pdf_dir = os.path.join(settings.MEDIA_ROOT, "expedientes_llenos", "pdf")
    os.makedirs(out_pdf_dir, exist_ok=True)

    nombre_base = f"Expediente_{s.folio}_{getattr(s, 'apellido_paterno_tutor', 'Usuario')}"
    nombre_base = _limpiar_nombre_archivo(nombre_base)
    out_pdf_path = os.path.join(out_pdf_dir, f"{nombre_base}.pdf")

    with open(out_pdf_path, "w+b") as result_file:
        pisa_status = pisa.CreatePDF(html, dest=result_file, link_callback=link_callback)

    if pisa_status.err:
        print(f"Error fatal al convertir Expediente a PDF HTML: {s.folio}")
        return None

    return out_pdf_path


# =========================================================
# SUPLENCIAS OPERATIVAS
# Pase de lista, autorizaciones temporales y bitácora.
#
# Estas funciones se mantienen en services.py para que views.py no concentre
# toda la lógica. No crean pantallas nuevas: solo entregan datos para que los
# paneles actuales muestren la viñeta de "Cubriendo temporalmente" cuando aplique.
# =========================================================


def _suplencias_models():
    """
    Importación perezosa para evitar ciclos entre models.py, views.py y services.py.
    """
    from admisiones.models import (
        AsistenciaPersonal,
        AutorizacionSuplencia,
        MovimientoSuplencia,
    )

    return AsistenciaPersonal, AutorizacionSuplencia, MovimientoSuplencia


def _usuario_model():
    from django.contrib.auth import get_user_model

    return get_user_model()


def _usuario_nombre_suplencia(usuario, default="Sin usuario"):
    """
    Nombre visible para pase de lista, autorizaciones y bitácoras.
    """
    if not usuario:
        return default

    for attr in ["nombre_visible", "nombre_completo"]:
        try:
            valor = getattr(usuario, attr, "")
            if callable(valor):
                valor = valor()
            if valor:
                return str(valor).strip()
        except Exception:
            pass

    try:
        nombre = usuario.get_full_name()
        if nombre:
            return nombre.strip()
    except Exception:
        pass

    return (
        getattr(usuario, "email", "")
        or getattr(usuario, "username", "")
        or str(usuario)
        or default
    )


def _normalizar_area_operativa(valor):
    """
    Normaliza áreas para comparar valores de usuarios, asistencia y suplencia.
    """
    valor = str(valor or "").strip().upper()
    valor = valor.replace("Á", "A").replace("É", "E").replace("Í", "I").replace("Ó", "O").replace("Ú", "U")
    valor = valor.replace(" ", "_").replace("-", "_")

    equivalencias = {
        "MEDICO": "MEDICA",
        "MEDICA": "MEDICA",
        "NUTRIOLOGA": "NUTRICION",
        "NUTRICIONAL": "NUTRICION",
        "NUTRICION": "NUTRICION",
        "PSICOLOGA": "PSICOLOGIA",
        "PSICOLOGICA": "PSICOLOGIA",
        "PSICOLOGIA": "PSICOLOGIA",
        "PEDAGOGIA": "PEDAGOGICA",
        "PEDAGOGICA": "PEDAGOGICA",
        "COORDINACION_DE_SALA": "PEDAGOGICA",
        "COORDINADORA_DE_SALA": "PEDAGOGICA",
        "TRABAJADORA_SOCIAL": "TRABAJO_SOCIAL",
        "TRABAJO_SOCIAL": "TRABAJO_SOCIAL",
        "DIRECCION": "DIRECCION",
        "DIRECTORA": "DIRECCION",
    }

    return equivalencias.get(valor, valor)


def _area_cubierta_desde_usuario(usuario):
    if not usuario:
        return ""

    return _normalizar_area_operativa(getattr(usuario, "area_operativa", ""))


def _area_display_suplencia(area):
    area = _normalizar_area_operativa(area)
    nombres = {
        "MEDICA": "Médica",
        "NUTRICION": "Nutrición",
        "PSICOLOGIA": "Psicología",
        "PEDAGOGICA": "Pedagógica",
        "TRABAJO_SOCIAL": "Trabajo Social",
        "DIRECCION": "Dirección",
    }
    return nombres.get(area, area.replace("_", " ").title() if area else "Sin área")


def usuario_es_direccion_suplencias(usuario):
    """
    Dirección es quien registra pase de lista y autoriza suplencias.
    También se permite superusuario/admin para soporte técnico.
    """
    if not usuario or not getattr(usuario, "is_authenticated", False):
        return False

    if bool(getattr(usuario, "is_superuser", False)):
        return True

    tipo_persona = str(getattr(usuario, "tipo_persona", "") or "").upper()
    area = _area_cubierta_desde_usuario(usuario)

    if area == "DIRECCION":
        return True

    if tipo_persona == "ADMIN":
        return True

    try:
        return usuario.groups.filter(name__icontains="directora").exists()
    except Exception:
        return False


def personal_operativo_pase_lista_queryset():
    """
    Personal que aparece en el pase de lista diario.

    Incluye únicamente:
    - Médica
    - Nutrióloga
    - Psicóloga
    - Pedagógica
    - Trabajadora Social

    No incluye Directora, Supervisora, Administrador, Beneficiarios ni Alumnos.
    """
    User = _usuario_model()

    try:
        if hasattr(User, "personal_para_pase_lista"):
            return User.personal_para_pase_lista()
    except Exception:
        pass

    return User.objects.filter(
        tipo_persona="PERSONAL",
        estatus_operativo=True,
        area_operativa__in=[
            "MEDICA",
            "NUTRICION",
            "PSICOLOGIA",
            "PEDAGOGICA",
            "TRABAJO_SOCIAL",
        ],
    ).order_by("area_operativa", "apellido_paterno", "apellido_materno", "nombre")


def personal_para_suplencias_queryset():
    """
    Personal que puede cubrir una ausencia.

    Incluye Dirección y personal operativo autorizado.
    No incluye Supervisora, Administrador del sistema, Beneficiarios ni Alumnos.
    """
    User = _usuario_model()

    try:
        if hasattr(User, "personal_para_suplencias"):
            return User.personal_para_suplencias()
    except Exception:
        pass

    return User.objects.filter(
        tipo_persona="PERSONAL",
        estatus_operativo=True,
        area_operativa__in=[
            "DIRECCION",
            "MEDICA",
            "NUTRICION",
            "PSICOLOGIA",
            "PEDAGOGICA",
            "TRABAJO_SOCIAL",
        ],
    ).order_by("area_operativa", "apellido_paterno", "apellido_materno", "nombre")


def titulares_para_suplencia_queryset():
    """
    Titulares que pueden marcarse como ausentes.
    Dirección no aparece aquí porque Dirección autoriza.
    """
    User = _usuario_model()

    try:
        if hasattr(User, "titulares_para_suplencia"):
            return User.titulares_para_suplencia()
    except Exception:
        pass

    return personal_operativo_pase_lista_queryset()


def asegurar_pase_lista_personal(fecha=None, semestre=None, registrado_por=None):
    """
    Crea registros NO_REGISTRADO para el pase de lista del día.
    Esto permite que Dirección vea a todo el personal operativo aun antes de marcar asistencia.
    """
    AsistenciaPersonal, _, _ = _suplencias_models()
    fecha = fecha or timezone.localdate()

    registros = []
    creados = 0

    for trabajador in personal_operativo_pase_lista_queryset():
        area = _area_cubierta_desde_usuario(trabajador)
        if not area:
            continue

        asistencia, created = AsistenciaPersonal.objects.get_or_create(
            fecha=fecha,
            trabajador=trabajador,
            defaults={
                "semestre": semestre,
                "area_operativa": area,
                "estado": AsistenciaPersonal.EstadoAsistencia.NO_REGISTRADO,
                "registrado_por": registrado_por,
            },
        )

        cambios = []
        if semestre and getattr(asistencia, "semestre_id", None) != getattr(semestre, "id", None):
            asistencia.semestre = semestre
            cambios.append("semestre")

        if area and asistencia.area_operativa != area:
            asistencia.area_operativa = area
            cambios.append("area_operativa")

        if cambios:
            asistencia.save(update_fields=cambios + ["actualizado_en"])

        if created:
            creados += 1

        registros.append(asistencia)

    return registros, creados


def guardar_pase_lista_personal(estados, usuario, fecha=None, semestre=None, observaciones=None):
    """
    Guarda el pase de lista capturado por Dirección.

    estados puede ser:
    - dict {trabajador_id: estado}
    - QueryDict con campos como estado_12, asistencia_12 o trabajador_12
    """
    from django.core.exceptions import ValidationError

    AsistenciaPersonal, _, _ = _suplencias_models()
    fecha = fecha or timezone.localdate()
    observaciones = observaciones or {}

    if not usuario_es_direccion_suplencias(usuario):
        return False, ["Solo Dirección puede registrar el pase de lista del personal."], []

    asegurar_pase_lista_personal(fecha=fecha, semestre=semestre, registrado_por=usuario)

    personal = list(personal_operativo_pase_lista_queryset())
    estados_validos = {opcion[0] for opcion in AsistenciaPersonal.EstadoAsistencia.choices}
    guardados = []
    errores = []

    def obtener_estado(trabajador_id):
        trabajador_id = str(trabajador_id)

        if isinstance(estados, dict):
            for clave in [
                trabajador_id,
                int(trabajador_id) if trabajador_id.isdigit() else trabajador_id,
                f"estado_{trabajador_id}",
                f"asistencia_{trabajador_id}",
                f"trabajador_{trabajador_id}",
            ]:
                if clave in estados:
                    valor = estados.get(clave)
                    if isinstance(valor, (list, tuple)):
                        valor = valor[0] if valor else ""
                    return str(valor or "").strip().upper()

        if hasattr(estados, "get"):
            return str(
                estados.get(f"estado_{trabajador_id}")
                or estados.get(f"asistencia_{trabajador_id}")
                or estados.get(f"trabajador_{trabajador_id}")
                or estados.get(trabajador_id)
                or ""
            ).strip().upper()

        return ""

    def obtener_observacion(trabajador_id):
        trabajador_id = str(trabajador_id)

        if isinstance(observaciones, dict):
            for clave in [trabajador_id, f"observaciones_{trabajador_id}", f"observacion_{trabajador_id}"]:
                if clave in observaciones:
                    return str(observaciones.get(clave) or "").strip()

        if hasattr(estados, "get"):
            return str(
                estados.get(f"observaciones_{trabajador_id}")
                or estados.get(f"observacion_{trabajador_id}")
                or ""
            ).strip()

        return ""

    for trabajador in personal:
        estado = obtener_estado(trabajador.id)
        if not estado:
            estado = AsistenciaPersonal.EstadoAsistencia.NO_REGISTRADO

        if estado not in estados_validos:
            errores.append(f"Estado no válido para {_usuario_nombre_suplencia(trabajador)}.")
            continue

        asistencia, _ = AsistenciaPersonal.objects.get_or_create(
            fecha=fecha,
            trabajador=trabajador,
            defaults={
                "semestre": semestre,
                "area_operativa": _area_cubierta_desde_usuario(trabajador),
                "estado": AsistenciaPersonal.EstadoAsistencia.NO_REGISTRADO,
                "registrado_por": usuario,
            },
        )

        asistencia.estado = estado
        asistencia.semestre = semestre or asistencia.semestre
        asistencia.area_operativa = _area_cubierta_desde_usuario(trabajador) or asistencia.area_operativa
        asistencia.actualizado_por = usuario
        asistencia.observaciones = obtener_observacion(trabajador.id)

        if not asistencia.registrado_por_id:
            asistencia.registrado_por = usuario

        try:
            asistencia.full_clean()
            asistencia.save()
            guardados.append(asistencia)
        except ValidationError as exc:
            errores.append(str(exc))
        except Exception as exc:
            errores.append(f"No se pudo guardar asistencia de {_usuario_nombre_suplencia(trabajador)}: {exc}")

    return not errores, errores, guardados


def estado_pase_lista_personal(fecha=None, semestre=None, crear_si_falta=True, usuario=None):
    """
    Devuelve el resumen del pase de lista para Dirección.
    """
    AsistenciaPersonal, _, _ = _suplencias_models()
    fecha = fecha or timezone.localdate()

    if crear_si_falta:
        asegurar_pase_lista_personal(fecha=fecha, semestre=semestre, registrado_por=usuario)

    personal = list(personal_operativo_pase_lista_queryset())
    asistencias = {
        asistencia.trabajador_id: asistencia
        for asistencia in AsistenciaPersonal.objects.filter(fecha=fecha).select_related("trabajador", "registrado_por")
    }

    filas = []
    total_presente = 0
    total_ausente = 0
    total_retardo = 0
    total_no_registrado = 0

    for trabajador in personal:
        asistencia = asistencias.get(trabajador.id)
        estado = getattr(asistencia, "estado", AsistenciaPersonal.EstadoAsistencia.NO_REGISTRADO)

        if estado == AsistenciaPersonal.EstadoAsistencia.PRESENTE:
            total_presente += 1
        elif estado in [
            AsistenciaPersonal.EstadoAsistencia.AUSENTE,
            AsistenciaPersonal.EstadoAsistencia.PERMISO,
            AsistenciaPersonal.EstadoAsistencia.INCAPACIDAD,
        ]:
            total_ausente += 1
        elif estado == AsistenciaPersonal.EstadoAsistencia.RETARDO:
            total_retardo += 1
        else:
            total_no_registrado += 1

        filas.append({
            "trabajador": trabajador,
            "trabajador_id": trabajador.id,
            "nombre": _usuario_nombre_suplencia(trabajador),
            "area_operativa": _area_cubierta_desde_usuario(trabajador),
            "area_display": _area_display_suplencia(_area_cubierta_desde_usuario(trabajador)),
            "asistencia": asistencia,
            "estado": estado,
            "estado_display": asistencia.get_estado_display() if asistencia else "No registrado",
            "permite_suplencia": bool(asistencia and asistencia.permite_suplencia),
            "esta_presente": estado == AsistenciaPersonal.EstadoAsistencia.PRESENTE,
            "sin_registrar": estado == AsistenciaPersonal.EstadoAsistencia.NO_REGISTRADO,
        })

    total_personal = len(personal)
    pase_lista_completo = bool(total_personal and total_no_registrado == 0)
    todo_personal_presente = bool(pase_lista_completo and total_ausente == 0)

    return {
        "fecha_pase_lista": fecha,
        "personal_pase_lista": personal,
        "asistencias_pase_lista": filas,
        "total_personal_pase_lista": total_personal,
        "total_presentes_pase_lista": total_presente,
        "total_ausentes_pase_lista": total_ausente,
        "total_retardos_pase_lista": total_retardo,
        "total_no_registrado_pase_lista": total_no_registrado,
        "pase_lista_completo": pase_lista_completo,
        "todo_personal_presente": todo_personal_presente,
        "hay_personal_ausente": total_ausente > 0,
        "ausencias_para_suplencia": [fila for fila in filas if fila["permite_suplencia"]],
    }


def actualizar_suplencias_vencidas(fecha=None):
    """
    Marca como vencidas las autorizaciones activas cuya vigencia ya terminó.
    """
    _, AutorizacionSuplencia, _ = _suplencias_models()
    ahora = timezone.now()
    qs = AutorizacionSuplencia.objects.filter(
        estado=AutorizacionSuplencia.EstadoAutorizacion.ACTIVA,
        fin__isnull=False,
        fin__lt=ahora,
    )

    if fecha:
        qs = qs.filter(fecha=fecha)

    return qs.update(estado=AutorizacionSuplencia.EstadoAutorizacion.VENCIDA, updated_at=ahora)


def obtener_suplencias_activas_usuario(usuario, fecha=None):
    _, AutorizacionSuplencia, _ = _suplencias_models()

    if not usuario or not getattr(usuario, "is_authenticated", False):
        return AutorizacionSuplencia.objects.none()

    actualizar_suplencias_vencidas(fecha=fecha)

    try:
        return AutorizacionSuplencia.activas_para_usuario(usuario, fecha=fecha).select_related(
            "titular_ausente",
            "suplente",
            "autorizada_por",
            "asistencia_personal",
        )
    except Exception:
        fecha = fecha or timezone.localdate()
        ahora = timezone.now()
        return AutorizacionSuplencia.objects.filter(
            suplente=usuario,
            fecha=fecha,
            estado=AutorizacionSuplencia.EstadoAutorizacion.ACTIVA,
            inicio__lte=ahora,
        ).filter(
            fin__isnull=True
        ) | AutorizacionSuplencia.objects.filter(
            suplente=usuario,
            fecha=fecha,
            estado=AutorizacionSuplencia.EstadoAutorizacion.ACTIVA,
            inicio__lte=ahora,
            fin__gte=ahora,
        )


def obtener_suplencia_activa_usuario_area(usuario, area_cubierta, accion=None, fecha=None):
    """
    Devuelve una autorización activa del usuario para el área y acción solicitadas.
    accion puede ser INCIDENCIA, DICTAMEN o TRABAJO_SOCIAL.
    """
    area_cubierta = _normalizar_area_operativa(area_cubierta)
    qs = obtener_suplencias_activas_usuario(usuario, fecha=fecha).filter(area_cubierta=area_cubierta)

    for autorizacion in qs:
        if not accion:
            return autorizacion
        try:
            if autorizacion.permite_accion(accion):
                return autorizacion
        except Exception:
            pass

    return None


def area_cubierta_desde_area_filtro(area_filtro):
    """
    Convierte nombres usados por filtros a áreas de suplencia.
    """
    area = _normalizar_area_operativa(area_filtro)
    mapeo = {
        "MEDICO": "MEDICA",
        "MEDICA": "MEDICA",
        "NUTRICION": "NUTRICION",
        "PSICOLOGIA": "PSICOLOGIA",
    }
    return mapeo.get(area, area)


def obtener_suplencia_activa_para_filtro(usuario, area_filtro, fecha=None):
    area = area_cubierta_desde_area_filtro(area_filtro)
    return obtener_suplencia_activa_usuario_area(usuario, area, accion="DICTAMEN", fecha=fecha)


def area_cubierta_desde_tipo_incidencia(tipo_incidencia):
    tipo = _normalizar_area_operativa(tipo_incidencia)
    mapeo = {
        "MEDICA": "MEDICA",
        "NUTRICIONAL": "NUTRICION",
        "NUTRICION": "NUTRICION",
        "PSICOLOGICA": "PSICOLOGIA",
        "PSICOLOGIA": "PSICOLOGIA",
        "PEDAGOGICA": "PEDAGOGICA",
    }
    return mapeo.get(tipo, tipo)


def obtener_suplencia_activa_para_tipo_incidencia(usuario, tipo_incidencia, fecha=None):
    area = area_cubierta_desde_tipo_incidencia(tipo_incidencia)
    return obtener_suplencia_activa_usuario_area(usuario, area, accion="INCIDENCIA", fecha=fecha)


def puede_autorizar_suplencia_para_asistencia(asistencia):
    """
    Reglas centrales:
    - Presente: no se puede suplir.
    - No registrado: primero debe registrarse pase de lista.
    - Ausente, permiso o incapacidad: sí permite suplencia.
    """
    AsistenciaPersonal, _, _ = _suplencias_models()

    if not asistencia:
        return False, "Primero debe registrarse el pase de lista del personal."

    if asistencia.estado == AsistenciaPersonal.EstadoAsistencia.PRESENTE:
        return False, "No se puede autorizar suplencia porque el personal titular está presente."

    if asistencia.estado == AsistenciaPersonal.EstadoAsistencia.NO_REGISTRADO:
        return False, "Primero debe completarse el pase de lista antes de autorizar suplencias."

    if not asistencia.permite_suplencia:
        return False, "Solo se permite suplencia cuando el estado sea Ausente, Permiso o Incapacidad."

    return True, "Suplencia permitida."


def crear_autorizacion_suplencia(
    titular_ausente,
    suplente,
    autorizada_por,
    asistencia_personal=None,
    tipo_cobertura=None,
    inicio=None,
    fin=None,
    motivo="",
    observaciones="",
    fecha=None,
):
    """
    Crea una autorización temporal de suplencia validando el pase de lista.
    """
    from django.core.exceptions import ValidationError

    AsistenciaPersonal, AutorizacionSuplencia, _ = _suplencias_models()
    fecha = fecha or timezone.localdate()
    inicio = inicio or timezone.now()

    if not usuario_es_direccion_suplencias(autorizada_por):
        return None, ["Solo Dirección puede autorizar suplencias."]

    area_cubierta = _area_cubierta_desde_usuario(titular_ausente)

    if not asistencia_personal:
        asistencia_personal = AsistenciaPersonal.objects.filter(
            fecha=fecha,
            trabajador=titular_ausente,
        ).first()

    ok_asistencia, mensaje_asistencia = puede_autorizar_suplencia_para_asistencia(asistencia_personal)
    if not ok_asistencia:
        return None, [mensaje_asistencia]

    if area_cubierta == AutorizacionSuplencia.AreaCubierta.TRABAJO_SOCIAL:
        tipo_cobertura = AutorizacionSuplencia.TipoCobertura.PROCESO_SOCIAL_COMPLETO
    elif not tipo_cobertura:
        tipo_cobertura = AutorizacionSuplencia.TipoCobertura.INCIDENCIA_DICTAMEN

    activa_existente = AutorizacionSuplencia.objects.filter(
        fecha=fecha,
        area_cubierta=area_cubierta,
        estado=AutorizacionSuplencia.EstadoAutorizacion.ACTIVA,
    ).first()

    if activa_existente:
        return None, [
            "Ya existe una suplencia activa para "
            f"{_area_display_suplencia(area_cubierta)} en esta fecha."
        ]

    autorizacion = AutorizacionSuplencia(
        fecha=fecha,
        asistencia_personal=asistencia_personal,
        titular_ausente=titular_ausente,
        suplente=suplente,
        autorizada_por=autorizada_por,
        area_cubierta=area_cubierta,
        tipo_cobertura=tipo_cobertura,
        inicio=inicio,
        fin=fin,
        motivo=motivo,
        observaciones=observaciones,
    )

    try:
        autorizacion.full_clean()
        autorizacion.save()
        return autorizacion, []
    except ValidationError as exc:
        if hasattr(exc, "message_dict"):
            errores = []
            for lista in exc.message_dict.values():
                errores.extend(lista if isinstance(lista, list) else [lista])
            return None, errores
        return None, [str(exc)]
    except Exception as exc:
        return None, [f"No se pudo crear la suplencia: {exc}"]


def cancelar_autorizacion_suplencia(autorizacion, usuario=None, motivo=""):
    if not autorizacion:
        return False, ["No se encontró la autorización de suplencia."]

    if not usuario_es_direccion_suplencias(usuario):
        return False, ["Solo Dirección puede cancelar suplencias."]

    try:
        return autorizacion.cancelar(usuario=usuario, motivo=motivo)
    except Exception as exc:
        return False, [str(exc)]


def registrar_movimiento_suplencia(
    autorizacion,
    usuario,
    accion=None,
    descripcion="",
    solicitud=None,
    inscripcion=None,
    incidencia=None,
    objeto=None,
    estado=None,
):
    """
    Guarda un movimiento en la bitácora de suplencias para que Supervisora lo consulte.
    """
    _, _, MovimientoSuplencia = _suplencias_models()

    if not autorizacion:
        return None

    if objeto is None:
        objeto = incidencia or inscripcion or solicitud

    if not accion:
        accion = MovimientoSuplencia.TipoAccion.OTRO

    if not descripcion:
        descripcion = "Movimiento registrado mediante suplencia autorizada por Dirección."

    objeto_modelo = ""
    objeto_id = None

    if objeto is not None:
        objeto_modelo = f"{objeto.__class__.__module__}.{objeto.__class__.__name__}"
        objeto_id = getattr(objeto, "id", None) or getattr(objeto, "pk", None)

    try:
        return MovimientoSuplencia.objects.create(
            autorizacion=autorizacion,
            usuario=usuario,
            titular_area=getattr(autorizacion, "titular_ausente", None),
            area_cubierta=getattr(autorizacion, "area_cubierta", ""),
            accion=accion,
            descripcion=descripcion,
            solicitud=solicitud,
            inscripcion=inscripcion,
            incidencia=incidencia,
            objeto_modelo=objeto_modelo,
            objeto_id=objeto_id,
            estado=estado or MovimientoSuplencia.EstadoMovimiento.PENDIENTE,
        )
    except Exception:
        return None


def marcar_dictamen_por_suplencia(solicitud, area_filtro, autorizacion, usuario=None):
    """
    Marca un dictamen como preliminar por suplencia si el modelo tiene los campos.
    No rompe el flujo si algún campo aún no existe.
    """
    if not solicitud or not autorizacion:
        return solicitud

    prefijos = {
        "PSICOLOGIA": "psicologia",
        "NUTRICION": "nutricion",
        "MEDICO": "medico",
        "MEDICA": "medico",
    }
    prefijo = prefijos.get(_normalizar_area_operativa(area_filtro))

    if not prefijo:
        return solicitud

    cambios = []

    asignaciones = {
        f"{prefijo}_es_suplencia": True,
        f"{prefijo}_autorizacion_suplencia": autorizacion,
        f"{prefijo}_suplente": usuario,
        f"{prefijo}_titular_ausente": getattr(autorizacion, "titular_ausente", None),
        f"{prefijo}_pendiente_revision_titular": True,
    }

    for campo, valor in asignaciones.items():
        if hasattr(solicitud, campo):
            setattr(solicitud, campo, valor)
            cambios.append(campo)

    if cambios:
        try:
            solicitud.save(update_fields=cambios)
        except Exception:
            solicitud.save()

    return solicitud


def marcar_incidencia_por_suplencia(incidencia, autorizacion, usuario=None):
    """
    Marca una incidencia como creada por suplencia si existen los campos en el modelo.
    """
    if not incidencia or not autorizacion:
        return incidencia

    cambios = []

    asignaciones = {
        "es_suplencia": True,
        "creada_por_suplencia": True,
        "autorizacion_suplencia": autorizacion,
        "area_original_suplencia": getattr(incidencia, "tipo", getattr(autorizacion, "area_cubierta", "")),
        "titular_ausente_suplencia": getattr(autorizacion, "titular_ausente", None),
        "suplente_suplencia": usuario or getattr(autorizacion, "suplente", None),
        "pendiente_revision_titular": True,
    }

    for campo, valor in asignaciones.items():
        if hasattr(incidencia, campo):
            setattr(incidencia, campo, valor)
            cambios.append(campo)

    if cambios:
        try:
            incidencia.save(update_fields=cambios)
        except Exception:
            incidencia.save()

    return incidencia


def marcar_trabajo_social_por_suplencia(inscripcion=None, solicitud=None, autorizacion=None, usuario=None, descripcion=""):
    """
    Marca movimientos del proceso de Trabajo Social hechos mediante cobertura completa.
    """
    _, AutorizacionSuplencia, MovimientoSuplencia = _suplencias_models()

    if not autorizacion and usuario:
        autorizacion = obtener_suplencia_activa_usuario_area(
            usuario,
            AutorizacionSuplencia.AreaCubierta.TRABAJO_SOCIAL,
            accion="TRABAJO_SOCIAL",
        )

    if not autorizacion:
        return None

    cambios = []

    if inscripcion:
        asignaciones = {
            "trabajo_social_es_suplencia": True,
            "trabajo_social_autorizacion_suplencia": autorizacion,
            "trabajo_social_suplente": usuario or getattr(autorizacion, "suplente", None),
            "trabajo_social_titular_ausente": getattr(autorizacion, "titular_ausente", None),
            "trabajo_social_pendiente_revision": True,
        }

        for campo, valor in asignaciones.items():
            if hasattr(inscripcion, campo):
                setattr(inscripcion, campo, valor)
                cambios.append(campo)

        if cambios:
            try:
                inscripcion.save(update_fields=cambios)
            except Exception:
                inscripcion.save()

    return registrar_movimiento_suplencia(
        autorizacion=autorizacion,
        usuario=usuario or getattr(autorizacion, "suplente", None),
        accion=MovimientoSuplencia.TipoAccion.PROCESO_TRABAJO_SOCIAL,
        descripcion=descripcion or "Movimiento de Trabajo Social realizado mediante cobertura temporal.",
        solicitud=solicitud,
        inscripcion=inscripcion,
        objeto=inscripcion or solicitud,
    )


def movimientos_suplencia_queryset():
    """
    Query base para la bitácora de suplencias que consulta Supervisora.
    """
    _, _, MovimientoSuplencia = _suplencias_models()

    return (
        MovimientoSuplencia.objects
        .select_related(
            "autorizacion",
            "autorizacion__titular_ausente",
            "autorizacion__suplente",
            "autorizacion__autorizada_por",
            "usuario",
            "titular_area",
            "solicitud",
            "inscripcion",
            "inscripcion__solicitud_ingreso",
            "incidencia",
        )
        .order_by("-creado_en")
    )


def filtrar_movimientos_suplencia(
    fecha_inicio=None,
    fecha_fin=None,
    area=None,
    estado=None,
    usuario=None,
    titular=None,
    texto=None,
):
    """
    Filtros simples para reportes de Supervisora.
    """
    from django.db.models import Q

    qs = movimientos_suplencia_queryset()

    if fecha_inicio:
        qs = qs.filter(creado_en__date__gte=fecha_inicio)

    if fecha_fin:
        qs = qs.filter(creado_en__date__lte=fecha_fin)

    if area:
        qs = qs.filter(area_cubierta=_normalizar_area_operativa(area))

    if estado:
        qs = qs.filter(estado=str(estado).strip().upper())

    if usuario:
        qs = qs.filter(usuario=usuario)

    if titular:
        qs = qs.filter(titular_area=titular)

    if texto:
        texto = str(texto).strip()
        qs = qs.filter(
            Q(descripcion__icontains=texto)
            | Q(solicitud__folio__icontains=texto)
            | Q(inscripcion__solicitud_ingreso__folio__icontains=texto)
            | Q(usuario__email__icontains=texto)
            | Q(titular_area__email__icontains=texto)
        )

    return qs


def contexto_suplencias_usuario(usuario, fecha=None):
    """
    Contexto para cualquier panel operativo: médica, nutrición, psicología,
    pedagógica o trabajo social.
    """
    suplencias = list(obtener_suplencias_activas_usuario(usuario, fecha=fecha))

    return {
        "suplencias_activas_usuario": suplencias,
        "tiene_suplencia_activa": bool(suplencias),
        "texto_suplencia_activa": "; ".join(
            f"Cubriendo temporalmente: {_area_display_suplencia(s.area_cubierta)}"
            for s in suplencias
        ),
    }


def contexto_suplencias_direccion(usuario=None, fecha=None, semestre=None, crear_pase=True):
    """
    Contexto para la viñeta de Dirección:
    - pase de lista
    - ausencias detectadas
    - suplencias activas
    - movimientos recientes
    """
    _, AutorizacionSuplencia, _ = _suplencias_models()
    fecha = fecha or timezone.localdate()

    estado = estado_pase_lista_personal(
        fecha=fecha,
        semestre=semestre,
        crear_si_falta=crear_pase,
        usuario=usuario,
    )

    actualizar_suplencias_vencidas(fecha=fecha)

    suplencias_activas = (
        AutorizacionSuplencia.objects
        .filter(fecha=fecha, estado=AutorizacionSuplencia.EstadoAutorizacion.ACTIVA)
        .select_related("titular_ausente", "suplente", "autorizada_por", "asistencia_personal")
        .order_by("area_cubierta", "inicio")
    )

    movimientos_recientes = movimientos_suplencia_queryset()[:10]

    return {
        **estado,
        "suplencias_activas_dia": suplencias_activas,
        "movimientos_suplencia_recientes": movimientos_recientes,
        "bitacora_suplencias": movimientos_recientes,
    }


def contexto_bitacora_suplencias(limite=100, **filtros):
    qs = filtrar_movimientos_suplencia(**filtros)

    return {
        "movimientos_suplencia": qs[:limite],
        "total_movimientos_suplencia": qs.count(),
    }


def expediente_permite_edicion_por_suplencia(inscripcion):
    """
    Trabajo Social por suplencia solo puede continuar expedientes abiertos.
    Finalizados, validados o cerrados quedan solo lectura.
    """
    if not inscripcion:
        return False

    estado = str(getattr(inscripcion, "estado", "") or "").strip().upper()

    return estado in {
        "NUEVA",
        "EN_PROCESO",
        "DEVUELTA",
        "EN_REVISION",
        "FILTROS",
        "REVISION_FINAL_TS",
        "ESPERANDO_DIRECTORA",
        "ACEPTADO_ESPERANDO_PAGO",
    }


def mensaje_bloqueo_expediente_suplencia(inscripcion):
    if expediente_permite_edicion_por_suplencia(inscripcion):
        return ""

    return (
        "Este expediente ya se encuentra finalizado, validado, rechazado o cerrado. "
        "En cobertura temporal de Trabajo Social solo puede visualizarse."
    )
